from js import Response, Headers
import json
import io

# We'll use python-pptx to create PowerPoint files
from pptx import Presentation
from pptx.util import Inches, Pt

async def on_fetch(request, env):
    """
    Main handler for the Python Worker
    This receives requests from the MCP Worker and creates PowerPoint files
    """
    
    # Handle CORS preflight
    if request.method == "OPTIONS":
        return Response.new("", {
            "headers": {
                "Access-Control-Allow-Origin": "*",
                "Access-Control-Allow-Methods": "POST, OPTIONS",
                "Access-Control-Allow-Headers": "Content-Type"
            }
        })
    
    # Only accept POST requests
    if request.method != "POST":
        return Response.new(
            json.dumps({"error": "Method not allowed. Use POST."}),
            {
                "status": 405,
                "headers": {"Content-Type": "application/json"}
            }
        )
    
    try:
        # Parse the incoming JSON request
        data = await request.json()
        
        # Validate required fields
        if not data:
            return Response.new(
                json.dumps({"error": "No data provided"}),
                {
                    "status": 400,
                    "headers": {"Content-Type": "application/json"}
                }
            )
        
        # Create the PowerPoint presentation
        prs = Presentation()
        
        # Get the title for the presentation
        presentation_title = data.get('title', 'My Presentation')
        
        # Add slides based on instructions
        slides_data = data.get('slides', [])
        
        if not slides_data:
            return Response.new(
                json.dumps({"error": "No slides provided. Please include 'slides' array in your request."}),
                {
                    "status": 400,
                    "headers": {"Content-Type": "application/json"}
                }
            )
        
        for slide_info in slides_data:
            layout_type = slide_info.get('layout', 'title_and_content')
            
            # Choose the right slide layout
            if layout_type == 'title':
                slide_layout = prs.slide_layouts[0]  # Title slide
            elif layout_type == 'title_and_content':
                slide_layout = prs.slide_layouts[1]  # Title and content
            elif layout_type == 'blank':
                slide_layout = prs.slide_layouts[6]  # Blank
            else:
                slide_layout = prs.slide_layouts[1]  # Default
            
            slide = prs.slides.add_slide(slide_layout)
            
            # Add title
            if 'title' in slide_info and slide.shapes.title:
                slide.shapes.title.text = slide_info['title']
            
            # Add bullets/content
            if 'bullets' in slide_info and len(slide.placeholders) > 1:
                body_shape = slide.placeholders[1]
                text_frame = body_shape.text_frame
                text_frame.clear()
                
                for i, bullet_text in enumerate(slide_info['bullets']):
                    if i == 0:
                        p = text_frame.paragraphs[0]
                    else:
                        p = text_frame.add_paragraph()
                    p.text = bullet_text
                    p.level = 0
        
        # Save to memory (not to disk)
        pptx_io = io.BytesIO()
        prs.save(pptx_io)
        pptx_bytes = pptx_io.getvalue()
        
        # Return the PowerPoint file
        headers = Headers.new({
            "Content-Type": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "Content-Disposition": f'attachment; filename="{presentation_title}.pptx"',
            "Access-Control-Allow-Origin": "*"
        }.items())
        
        return Response.new(pptx_bytes, {"headers": headers})
    
    except json.JSONDecodeError:
        return Response.new(
            json.dumps({"error": "Invalid JSON in request body"}),
            {
                "status": 400,
                "headers": {"Content-Type": "application/json"}
            }
        )
    except Exception as e:
        return Response.new(
            json.dumps({"error": f"Failed to create presentation: {str(e)}"}),
            {
                "status": 500,
                "headers": {"Content-Type": "application/json"}
            }
        )
